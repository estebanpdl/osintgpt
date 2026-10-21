'''Text identity, locations, and cache-only reads without paid providers.'''

import json
from dataclasses import replace

import pytest

from osintgpt.ingestion import Document, FieldMapping
from osintgpt.ingestion.provenance import quote_locations, text_spans
from osintgpt.ingestion.snapshots import load_snapshot, read_snapshot
from osintgpt.ingestion.transcription import cache_path


def test_page_boundaries_preserve_sentence_join_and_exact_offsets(tmp_path, monkeypatch):
    path = tmp_path / 'paper.pdf'
    path.write_bytes(b'pdf fixture')
    pages = ['A long introduction describes how Alpha funded',
             'the Beta operation in May, according to this report.']
    monkeypatch.setattr('osintgpt.ingestion.pdf.extract_page_texts', lambda path: pages)
    snapshot = load_snapshot(path, cache_dir=tmp_path / 'extracts')
    document = snapshot.documents[0]
    matches = quote_locations(document, 'Alpha funded the Beta operation')
    assert {span['page'] for span in matches[0]['spans']} == {1, 2}
    assert document.text[matches[0]['start']:matches[0]['end']] == 'Alpha funded the Beta operation'


def test_cache_only_scan_read_reuses_existing_transcription(tmp_path, monkeypatch):
    path, cache = tmp_path / 'scan.pdf', tmp_path / 'extracts'
    path.write_bytes(b'scanned pdf')
    cache.mkdir()
    cache_path(cache, b'page image').write_text('Alpha funds Beta.', encoding='utf-8')
    monkeypatch.setattr('osintgpt.ingestion.pdf.extract_page_texts', lambda path: [''])
    monkeypatch.setattr('osintgpt.ingestion.pdf.render_page', lambda *args: b'page image')
    snapshot = load_snapshot(path, cache_dir=cache)
    assert snapshot.documents[0].text == 'Alpha funds Beta.'
    assert snapshot.documents[0].spans[0].transcribed
    assert snapshot.missing_pages == []


def test_gap_is_explicit_and_later_transcription_creates_new_version(tmp_path, monkeypatch):
    path, cache = tmp_path / 'scan.pdf', tmp_path / 'extracts'
    path.write_bytes(b'scanned pdf')
    monkeypatch.setattr('osintgpt.ingestion.pdf.extract_page_texts', lambda path: [''])
    monkeypatch.setattr('osintgpt.ingestion.pdf.render_page', lambda *args: b'page')
    gap = load_snapshot(path, cache_dir=cache)
    assert gap.missing_pages == [1]
    assert quote_locations(gap.documents[0], 'no extractable text') == []
    complete = load_snapshot(path, cache_dir=cache, transcriber=lambda *args: 'Alpha funds Beta.')
    assert complete.fingerprint == gap.fingerprint
    assert complete.version != gap.version and complete.missing_pages == []
    assert read_snapshot(cache, gap.version).missing_pages == [1]


def test_complete_snapshot_reused_without_loader_or_transcriber(tmp_path, monkeypatch):
    path, cache = tmp_path / 'a.md', tmp_path / 'extracts'
    path.write_text('# Section\n\nAlpha funds Beta.')
    first = load_snapshot(path, ref='a.md', cache_dir=cache)
    def forbidden(*args, **kwargs):
        raise AssertionError('A cached snapshot must not extract or transcribe again')
    monkeypatch.setattr('osintgpt.ingestion.snapshots.load_documents', forbidden)
    assert load_snapshot(path, ref='a.md', cache_dir=cache, transcriber=forbidden) == first


def test_source_and_mapping_changes_preserve_immutable_prior_text(tmp_path):
    path, cache = tmp_path / 'records.csv', tmp_path / 'extracts'
    path.write_text('id,body,title\na,Alpha funds Beta.,Report\n')
    mapping = FieldMapping(content=('body',), identity='id')
    first = load_snapshot(path, mapping, ref='records.csv', cache_dir=cache)
    second = load_snapshot(path, replace(mapping, content=('title',)), ref='records.csv', cache_dir=cache)
    path.write_text('id,body,title\na,Alpha stopped funding Beta.,Report\n')
    third = load_snapshot(path, mapping, ref='records.csv', cache_dir=cache)
    assert len({item.version for item in (first, second, third)}) == 3
    assert read_snapshot(cache, first.version).documents[0].text == 'Alpha funds Beta.'
    assert first.documents[0].ref == third.documents[0].ref == 'records.csv#a'


@pytest.mark.parametrize('suffix', ['.csv', '.xlsx', '.json', '.jsonl'])
def test_structured_records_retain_distinct_identity(tmp_path, suffix):
    path = tmp_path / ('records' + suffix)
    rows = [{'id': 'a', 'body': 'Alpha funds Beta.'}, {'id': 'b', 'body': 'Gamma funds Delta.'}]
    if suffix == '.csv':
        path.write_text('id,body\na,Alpha funds Beta.\nb,Gamma funds Delta.\n')
    elif suffix == '.xlsx':
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(['id', 'body'])
        for row in rows:
            workbook.active.append(list(row.values()))
        workbook.save(path)
        workbook.close()
    else:
        path.write_text(json.dumps(rows) if suffix == '.json' else '\n'.join(map(json.dumps, rows)))
    snapshot = load_snapshot(path, FieldMapping(content=('body',), identity='id'), ref=path.name)
    assert [document.ref for document in snapshot.documents] == [path.name + '#a', path.name + '#b']
    assert snapshot.documents[0].text == rows[0]['body']
    assert snapshot.documents[1].text == rows[1]['body']


def test_duplicate_record_ids_fail_instead_of_merging_evidence(tmp_path):
    path = tmp_path / 'rows.csv'
    path.write_text('id,body\na,One\na,Two\n')
    with pytest.raises(ValueError, match='Duplicate record identities'):
        load_snapshot(path, FieldMapping(content=('body',), identity='id'))


def test_all_quote_occurrences_and_nfc_offsets_are_preserved():
    text = '# First\n\nCafe\u0301\n supports  Beta.\n\n# Second\n\nCafé supports Beta.'
    document = Document('a.md', text, spans=text_spans(text))
    # Normalize the entire source even when one occurrence already matches exactly.
    matches = quote_locations(document, 'Café supports Beta.')
    assert len(matches) == 2
    assert [item['spans'][0]['section'] for item in matches] == ['First', 'Second']
    assert text[matches[0]['start']:matches[0]['end']] == 'Cafe\u0301\n supports  Beta.'


def test_hangul_normalization_preserves_original_offsets():
    import unicodedata
    quote = '한국어 문장'
    text = unicodedata.normalize('NFD', quote)
    document = Document('a.md', text, spans=text_spans(text))
    match = quote_locations(document, quote)[0]
    assert text[match['start']:match['end']] == text


def test_snapshot_damage_and_path_traversal_are_rejected(tmp_path):
    path, cache = tmp_path / 'a.md', tmp_path / 'extracts'
    path.write_text('Alpha funds Beta.')
    snapshot = load_snapshot(path, cache_dir=cache)
    saved = cache / 'text' / f'{snapshot.version}.json'
    saved.write_text(saved.read_text().replace('Alpha', 'Gamma'))
    with pytest.raises(ValueError, match='damaged'):
        read_snapshot(cache, snapshot.version)
    with pytest.raises(ValueError, match='Invalid'):
        read_snapshot(cache, '../secret')


def test_source_change_during_extraction_is_not_cached(tmp_path, monkeypatch):
    path, cache = tmp_path / 'a.md', tmp_path / 'extracts'
    path.write_text('Before')
    def changing(*args, **kwargs):
        path.write_text('After')
        return [Document(path.as_posix(), 'Before')]
    monkeypatch.setattr('osintgpt.ingestion.snapshots.load_documents', changing)
    with pytest.raises(ValueError, match='Source changed'):
        load_snapshot(path, cache_dir=cache)
    assert not (cache / 'text').exists()


def test_correcting_cached_transcription_versions_text_without_new_vision_call(tmp_path, monkeypatch):
    path, cache = tmp_path / 'scan.pdf', tmp_path / 'extracts'
    path.write_bytes(b'scan')
    cache.mkdir()
    transcript = cache_path(cache, b'page')
    transcript.write_text('Alpha funds Beta.')
    monkeypatch.setattr('osintgpt.ingestion.pdf.extract_page_texts', lambda path: [''])
    monkeypatch.setattr('osintgpt.ingestion.pdf.render_page', lambda *args: b'page')
    before = load_snapshot(path, cache_dir=cache)
    transcript.write_text('Alpha funds Gamma.')
    after = load_snapshot(path, cache_dir=cache)
    assert before.fingerprint == after.fingerprint and before.version != after.version
    assert read_snapshot(cache, before.version).documents[0].text == 'Alpha funds Beta.'
    assert after.documents[0].text == 'Alpha funds Gamma.'


def test_generated_snapshots_are_not_discovered_as_new_corpus(tmp_path):
    from osintgpt.ingestion.sources import Source
    path = tmp_path / 'a.md'
    path.write_text('Alpha funds Beta.')
    load_snapshot(path, cache_dir=tmp_path / 'extracts')
    assert Source('.').resolve(tmp_path) == [path]


def test_thousand_page_text_retains_last_page_location(tmp_path, monkeypatch):
    path = tmp_path / 'large.pdf'
    path.write_bytes(b'pdf fixture')
    pages = [f'Page {number} describes a distinct transaction from Alpha to Beta.'
             for number in range(1, 1001)]
    monkeypatch.setattr('osintgpt.ingestion.pdf.extract_page_texts', lambda path: pages)
    snapshot = load_snapshot(path, ref='large.pdf', cache_dir=tmp_path / 'extracts')
    matches = quote_locations(snapshot.documents[0], pages[-1])
    assert len(matches) == 1 and matches[0]['spans'][0]['page'] == 1000
