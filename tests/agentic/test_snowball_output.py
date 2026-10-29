'''Lossless rows, spreadsheet safety and date grouping contracts.'''

import csv
import io
import json
from dataclasses import replace

import pytest
from openpyxl import load_workbook

from osintgpt import snowball
from osintgpt.agentic.snowball_output import (
    EXPORT_FIELDS, export_snowball, parse_date, snowball_csv, snowball_data,
    snowball_rows, snowball_xlsx,
)
from tests.snowball_fakes import GraphEmbedder, GraphStore


def test_csv_and_xlsx_have_identical_columns_full_text_and_copy_sources(tmp_path):
    result = snowball(None, 'start', GraphEmbedder(), store=GraphStore(4, 8), depth=4)
    expected = snowball_rows(result, 'month')
    csv_rows = list(csv.DictReader(io.StringIO(snowball_csv(result, 'month').lstrip('\ufeff'))))
    assert tuple(csv_rows[0]) == EXPORT_FIELDS
    assert [row['text'] for row in csv_rows] == [row['text'] for row in expected]
    assert [int(row['hop']) for row in csv_rows] == [1, 2, 3, 4]
    assert json.loads(csv_rows[0]['identical_text_refs']) == result.hops[0].identical_text_refs
    assert 'similarity_to_query' in csv_rows[0] and 'drift' not in csv_rows[0]
    book = load_workbook(io.BytesIO(snowball_xlsx(result, 'month')))
    rows = list(book.active.values)
    assert rows[0] == EXPORT_FIELDS
    assert rows[1][EXPORT_FIELDS.index('score')] == pytest.approx(expected[0]['score'])
    assert rows[3][EXPORT_FIELDS.index('period')] == 'Unknown date'
    assert rows[1][EXPORT_FIELDS.index('identical_text_refs')] == csv_rows[0]['identical_text_refs']
    book.close()
    for suffix in ('csv', 'xlsx'):
        path = tmp_path / f'chain.{suffix}'
        export_snowball(result, path)
        assert path.stat().st_size > 0
        with pytest.raises(FileExistsError):
            export_snowball(result, path)


def test_formula_like_strings_are_reversibly_escaped_but_scores_stay_numeric():
    result = snowball(None, 'start', GraphEmbedder(), store=GraphStore(), depth=1)
    hop = result.hops[0]
    hit = replace(hop.result, chunk=replace(hop.result.chunk, text='=HYPERLINK("fake")'), score=-.2)
    result = replace(result, hops=[replace(hop, query=' +SUM(A1)', result=hit)])
    row = next(csv.DictReader(io.StringIO(snowball_csv(result).lstrip('\ufeff'))))
    assert row['query'].startswith("' +") and row['text'].startswith("'=")
    assert set(json.loads(row['spreadsheet_escaped_fields'])) == {'query', 'text'}
    assert float(row['score']) == -.2
    book = load_workbook(io.BytesIO(snowball_xlsx(result)))
    assert book.active.cell(2, 3).data_type == 's'
    assert book.active.cell(2, 8).value == -.2
    book.close()


@pytest.mark.parametrize('text', ['x' * 32768, 'bad\x00control'], ids=['long-cell', 'control-character'])
def test_xlsx_rejects_silent_loss_and_csv_keeps_the_text(text):
    result = snowball(None, 'start', GraphEmbedder(), store=GraphStore(), depth=1)
    hop = result.hops[0]
    hit = replace(hop.result, chunk=replace(hop.result.chunk, text=text))
    result = replace(result, hops=[replace(hop, result=hit)])
    with pytest.raises(ValueError, match='losslessly'):
        snowball_xlsx(result)
    assert text in snowball_csv(result)


@pytest.mark.parametrize(('text', 'serialized'), [
    ('plain\\path', 'plain\\path'),
    ('a,b', '"a,b"'),
    ('say "hi"', '"say ""hi"""'),
    ('line\nbreak', '"line\nbreak"'),
    ('line\r\nbreak', '"line\r\nbreak"'),
], ids=['backslash', 'comma', 'quote', 'newline', 'crlf'])
def test_csv_keeps_standard_quoting_and_roundtrips_text(text, serialized):
    result = snowball(None, 'start', GraphEmbedder(), store=GraphStore(), depth=1)
    hop = result.hops[0]
    hit = replace(hop.result, chunk=replace(hop.result.chunk, text=text))
    result = replace(result, hops=[replace(hop, result=hit)])
    exported = snowball_csv(result)
    assert f',{serialized},' in exported
    row = next(csv.DictReader(io.StringIO(exported.lstrip('\ufeff'), newline='')))
    assert row['text'] == text


def test_empty_exports_still_have_columns_and_json_is_serializable():
    result = snowball(None, 'start', GraphEmbedder(), store=GraphStore(0))
    assert snowball_csv(result).lstrip('\ufeff').strip() == ','.join(EXPORT_FIELDS)
    book = load_workbook(io.BytesIO(snowball_xlsx(result)))
    assert list(book.active.values) == [EXPORT_FIELDS]
    book.close()
    assert json.loads(json.dumps(snowball_data(result)))['hop_count'] == 0


def test_dates_keep_timezone_and_iso_week_year():
    parsed = parse_date('2021-01-01T23:00:00-05:00')
    assert parsed.isoformat() == '2021-01-01T23:00:00-05:00'
    result = snowball(None, 'start', GraphEmbedder(), store=GraphStore(), depth=1)
    result = replace(result, hops=[replace(result.hops[0], date=parsed.isoformat())])
    assert snowball_rows(result, 'week')[0]['period'] == '2020-W53'
    assert parse_date('03/04/2026') is None
    assert parse_date('03/04/2026', '%d/%m/%Y').month == 4


def test_invalid_grouping_and_export_suffix_are_explicit(tmp_path):
    result = snowball(None, 'start', GraphEmbedder(), store=GraphStore(), depth=1)
    with pytest.raises(ValueError, match='Group by'):
        snowball_rows(result, 'hour')
    with pytest.raises(ValueError, match='.csv or .xlsx'):
        export_snowball(result, tmp_path / 'out.txt')
